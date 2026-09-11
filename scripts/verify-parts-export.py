"""Independent read-only QA of browser-exported OOXML with openpyxl."""
import datetime
import json
import pathlib
import xml.etree.ElementTree as ET
import zipfile
import openpyxl

root=pathlib.Path(__file__).resolve().parents[1]
original=root/'public/parts-workbook-template.xlsx'
exported=root/'work/parts-export-qa/export-original.xlsx'
edited=root/'work/parts-export-qa/export-edited.xlsx'
source=openpyxl.load_workbook(original)
result=openpyxl.load_workbook(exported)
cached_source=openpyxl.load_workbook(original,data_only=True)
cached_result=openpyxl.load_workbook(exported,data_only=True)
count=0
for before,after in zip(source,result):
    assert before.title==after.title
    assert before.max_row==after.max_row and before.max_column==after.max_column
    assert set(map(str,before.merged_cells.ranges))==set(map(str,after.merged_cells.ranges))
    # The export resets scrolling to the top, retaining the frozen split itself.
    a,b=before.sheet_view.pane,after.sheet_view.pane
    assert (a is None)==(b is None)
    if a is not None:
        assert (a.xSplit,a.ySplit,a.state)==(b.xSplit,b.ySplit,b.state)
    assert dict(before.page_setup)==dict(after.page_setup)
    assert dict(before.page_margins)==dict(after.page_margins)
    for key,dimension in before.column_dimensions.items():
        assert dict(dimension)==dict(after.column_dimensions[key]),('column',before.title,key)
    for row in before:
        for cell in row:
            found=after[cell.coordinate]
            assert cell.value==found.value,(before.title,cell.coordinate,cell.value,found.value)
            assert cell._style==found._style,('style',before.title,cell.coordinate)
            assert cached_source[before.title][cell.coordinate].value==cached_result[before.title][cell.coordinate].value,('cached',before.title,cell.coordinate)
            count+=1
    for key,dimension in before.row_dimensions.items():
        assert dict(dimension)==dict(after.row_dimensions[key]),('row',before.title,key)

changed=openpyxl.load_workbook(edited)
cached_changed=openpyxl.load_workbook(edited,data_only=True)
assert changed['Sheet1']['H5'].value=='NEW-PART'
assert changed['Sheet1']['J5'].value==12
assert cached_changed['Sheet1']['J4'].value==30
assert cached_changed['Sheet1']['T4'].value==0
assert cached_changed['Sheet1']['U4'].value==30
assert cached_changed['Sheet1']['N4'].value==datetime.datetime(2026,9,11)
assert changed['Sheet1']['AA4'].value=='导出验证 & <保留格式>'
assert changed['Sheet1']['H4']._style==source['Sheet1']['I4']._style
assert changed['Sheet1'].column_dimensions['H'].width==source['Sheet1'].column_dimensions['I'].width
assert changed['Sheet1']['E4'].value=='=VLOOKUP(H:H,[1]Sheet1!$J:$P,7,0)'
with zipfile.ZipFile(exported) as archive:
    for name in archive.namelist():
        if name.endswith('.xml') or name.endswith('.rels'):
            ET.fromstring(archive.read(name))
print(json.dumps({'verified_cells':count,'styles':'identical','values_and_formulas':'identical','print_settings':'identical','edited_export':'passed'}))
